"""
複数ファイルの差分を1つのExcelサマリブックに集約するレンダラー。

`html_renderer.render(file_diff) -> str` と対になる純粋関数群。
HTMLをパースするのではなく、`FileDiff` のリストを直接入力として
openpyxl の `Workbook` を組み立てる。
"""
from __future__ import annotations

from difflib import SequenceMatcher
from typing import Optional

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .diff_engine import FileDiff, RowTag, SheetDiff
from .html_renderer import _strip_ctrl
from .reader import CellData, RowData

# ---------------------------------------------------------------------------
# 列構成・色定義
# ---------------------------------------------------------------------------

_HEADERS = [
    "旧ファイル", "新ファイル", "シート名", "種類",
    "比較キー", "準比較キー", "項目", "旧項目値", "新項目値",
]
_COL_WIDTHS = [40, 40, 16, 10, 14, 14, 18, 40, 40]

# HTML版 (html_renderer.py) の配色を流用し、HTMLとExcelで見た目の一貫性を保つ。
# 変更（MODIFY）行は行全体の背景色を付けない（うるさいためユーザー確定済み）。
# H/I列のリッチテキスト色分けと、サブキー救済ペアの紫背景のみで変更箇所を示す。
_FILL_ADD = PatternFill("solid", fgColor="FFE6FFED")     # .row-inserted td
_FILL_DEL = PatternFill("solid", fgColor="FFFFEEF0")     # .row-deleted td
_FILL_SUBKEY = PatternFill("solid", fgColor="FFE8DCFF")  # .cell-modified-subkey
_FILL_HEADER = PatternFill("solid", fgColor="FF1F3864")

_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_DEL_COLOR = "CF222E"  # 赤字（H列：削除された文字部分）
_INS_COLOR = "1A7F37"  # 緑字（I列：追加された文字部分）
_TOP_LEFT = Alignment(horizontal="left", vertical="top", wrap_text=False)


# ---------------------------------------------------------------------------
# ヘルパー: 比較キー・項目名
# ---------------------------------------------------------------------------

def _sanitize_xml_text(s: str) -> str:
    """XML 1.0で許可されない制御文字を除去する。

    プレーンなセル値は openpyxl が書き込み時に ILLEGAL_CHARACTERS_RE で検証し
    IllegalCharacterError を送出するが、CellRichText/TextBlock 経由の値は
    この検証を通らず、そのまま壊れたXMLとして書き込まれてしまう
    （Excel起動時に「修復されたレコード」として警告が出る）。
    そのため本レンダラーでは全ての出力テキストに対して明示的に適用する。
    """
    return ILLEGAL_CHARACTERS_RE.sub("", s)


def _cell_display(row: Optional[RowData], col_idx: int) -> str:
    if row is None or col_idx >= len(row.cells):
        return ""
    return _sanitize_xml_text(row.cells[col_idx].display())


def _key_value_str(row: Optional[RowData], cols: list[int]) -> str:
    """key_cols/sub_key_cols で指定された列の値をカンマ結合して返す。"""
    if not cols or row is None:
        return ""
    return ",".join(_cell_display(row, c) for c in cols)


def _resolve_header_labels(
    sheet_diff: SheetDiff,
    header_row: int,
) -> dict[int, str]:
    """row_diffs から header_row 行の値を集め、col_idx -> ラベル の辞書を返す。

    old側の値を優先し、空ならnew側を使う。値が解決できない列はこの辞書に
    含めない（呼び出し側の _header_label() で列番号表記にフォールバックする）。
    """
    if header_row <= 0:
        return {}

    old_header: Optional[RowData] = None
    new_header: Optional[RowData] = None
    for rd in sheet_diff.row_diffs:
        if (old_header is None and rd.old_row is not None
                and rd.old_row.row_idx == header_row):
            old_header = rd.old_row
        if (new_header is None and rd.new_row is not None
                and rd.new_row.row_idx == header_row):
            new_header = rd.new_row
        if old_header is not None and new_header is not None:
            break

    labels: dict[int, str] = {}
    for col_idx in range(sheet_diff.max_cols):
        val = _cell_display(old_header, col_idx) or _cell_display(new_header, col_idx)
        if val:
            labels[col_idx] = val
    return labels


def _header_label(labels: dict[int, str], col_idx: int) -> str:
    return labels.get(col_idx) or get_column_letter(col_idx + 1)


# ---------------------------------------------------------------------------
# ヘルパー: 文字単位diffのリッチテキスト化
# ---------------------------------------------------------------------------

def _char_diff_runs(
    old_cell: Optional[CellData],
    new_cell: Optional[CellData],
) -> tuple[list[tuple[str, bool]], list[tuple[str, bool]]]:
    """(old側runs, new側runs) を返す。各runは (text, is_highlighted)。

    SequenceMatcher のopcodesロジックは html_renderer._render_cell_pair_diff
    と同一（移植元）。
    """
    old_str = _sanitize_xml_text(_strip_ctrl(old_cell.value) if old_cell else "")
    new_str = _sanitize_xml_text(_strip_ctrl(new_cell.value) if new_cell else "")

    old_runs: list[tuple[str, bool]] = []
    new_runs: list[tuple[str, bool]] = []

    matcher = SequenceMatcher(None, old_str, new_str, autojunk=False)
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        oc = old_str[i1:i2]
        nc = new_str[j1:j2]
        if tag == "equal":
            if oc:
                old_runs.append((oc, False))
            if nc:
                new_runs.append((nc, False))
        elif tag == "delete":
            old_runs.append((oc, True))
        elif tag == "insert":
            new_runs.append((nc, True))
        elif tag == "replace":
            old_runs.append((oc, True))
            new_runs.append((nc, True))

    return old_runs, new_runs


def _build_rich_value(
    runs: list[tuple[str, bool]],
    highlight_color: str,
    highlight_bold: bool,
    highlight_strike: bool,
    whole_cell_strike: bool,
):
    """runs から CellRichText（必要なら）またはプレーン文字列を組み立てて返す。

    highlight_* は「変更部分」のrunに適用するフォント指定。
    whole_cell_strike は CellData.strikethrough によるセル全体への取り消し線
    （--strikethrough 指定時のみ意味を持つ）で、全runに追加適用される。
    """
    if not runs:
        return ""

    needs_rich = whole_cell_strike or any(is_hl for _, is_hl in runs)
    if not needs_rich:
        return "".join(text for text, _ in runs)

    blocks: list = []
    for text, is_hl in runs:
        if is_hl:
            font = InlineFont(
                color=highlight_color,
                b=highlight_bold or None,
                strike=(highlight_strike or whole_cell_strike) or None,
            )
        elif whole_cell_strike:
            font = InlineFont(strike=True)
        else:
            font = None
        blocks.append(TextBlock(font, text) if font is not None else text)

    return CellRichText(blocks)


def _render_cell_pair_rich(
    old_cell: Optional[CellData],
    new_cell: Optional[CellData],
) -> tuple:
    """変更セルペアから (H列値, I列値) を返す。値はプレーン文字列 or CellRichText。"""
    old_runs, new_runs = _char_diff_runs(old_cell, new_cell)
    old_strike = bool(old_cell and old_cell.strikethrough)
    new_strike = bool(new_cell and new_cell.strikethrough)
    old_val = _build_rich_value(
        old_runs, _DEL_COLOR, highlight_bold=False, highlight_strike=True,
        whole_cell_strike=old_strike,
    )
    new_val = _build_rich_value(
        new_runs, _INS_COLOR, highlight_bold=True, highlight_strike=False,
        whole_cell_strike=new_strike,
    )
    return old_val, new_val


# ---------------------------------------------------------------------------
# 行書き込み
# ---------------------------------------------------------------------------

def _write_row(
    ws: Worksheet,
    row_idx: int,
    values: list,
    fill: Optional[PatternFill],
    subkey: bool = False,
) -> None:
    """A〜I列に1行書き込む。fill が None でなければ行全体を塗る
    （subkey=True の場合はH/I列のみ紫で上書きし、行全体は塗らない）。
    """
    for col_idx, val in enumerate(values, start=1):
        if isinstance(val, str):
            # 生の \r / \r\n はXML往復時にパーサーが暗黙的に \n へ正規化してしまい
            # （OOXMLのST_Xstring往復仕様上は数値文字参照 &#13; でエスケープすべき）、
            # Excelの検証で「修復」対象になるため、他の列と同様にLFへ正規化する。
            val = _sanitize_xml_text(_strip_ctrl(val))
        c = ws.cell(row=row_idx, column=col_idx, value=(val if val != "" else None))
        c.alignment = _TOP_LEFT
        if subkey:
            if col_idx in (8, 9):  # H, I
                c.fill = _FILL_SUBKEY
        elif fill is not None:
            c.fill = fill


# ---------------------------------------------------------------------------
# 公開関数
# ---------------------------------------------------------------------------

def render(
    file_diffs: list[FileDiff],
    header_row: int = 1,
    sub_key_cols: Optional[list[int]] = None,
) -> Workbook:
    """FileDiff のリストを受け取り、集約Excelの Workbook を返す。

    単一ファイル比較の場合も「要素数1のリスト」として同じコードパスに乗せる。
    """
    sub_key_cols = sub_key_cols or []

    wb = Workbook()
    ws = wb.active
    ws.title = "差分一覧"

    for col_idx, hdr in enumerate(_HEADERS, start=1):
        c = ws.cell(row=1, column=col_idx, value=hdr)
        c.font = _HEADER_FONT
        c.fill = _FILL_HEADER
    ws.freeze_panes = "A2"

    row_idx = 2
    for file_diff in file_diffs:
        old_path = file_diff.old_path
        new_path = file_diff.new_path

        for sheet_diff in file_diff.sheet_diffs:
            if sheet_diff.status == "equal":
                continue

            if sheet_diff.status in ("added", "deleted"):
                is_added = sheet_diff.status == "added"
                kind = "シート追加" if is_added else "シート削除"
                fill = _FILL_ADD if is_added else _FILL_DEL
                _write_row(
                    ws, row_idx,
                    [old_path, new_path, sheet_diff.name, kind, "", "", "", "", ""],
                    fill,
                )
                row_idx += 1
                continue

            # status == "modified"
            header_labels = _resolve_header_labels(sheet_diff, header_row)
            key_cols = sheet_diff.key_cols

            for rd in sheet_diff.row_diffs:
                if rd.tag == RowTag.EQUAL:
                    continue

                if rd.tag in (RowTag.DELETE, RowTag.INSERT):
                    is_delete = rd.tag == RowTag.DELETE
                    kind = "削除" if is_delete else "追加"
                    fill = _FILL_DEL if is_delete else _FILL_ADD
                    src_row = rd.old_row if is_delete else rd.new_row
                    key_val = _key_value_str(src_row, key_cols)
                    sub_key_val = _key_value_str(src_row, sub_key_cols)
                    _write_row(
                        ws, row_idx,
                        [old_path, new_path, sheet_diff.name, kind,
                         key_val, sub_key_val, "", "", ""],
                        fill,
                    )
                    row_idx += 1
                    continue

                # RowTag.MODIFY: 1 CellDiff = 1行
                key_val = (
                    _key_value_str(rd.old_row, key_cols)
                    or _key_value_str(rd.new_row, key_cols)
                )
                # 準比較キー（F列）は --sub-key-cols 指定時、サブキー救済の有無に
                # 関わらず全ての変更行に値を入れる（ユーザー確定済み）。
                sub_key_val = (
                    _key_value_str(rd.old_row, sub_key_cols)
                    or _key_value_str(rd.new_row, sub_key_cols)
                )
                is_subkey_row = rd.matched_by == "subkey"

                for cd in rd.cell_diffs:
                    item_label = _header_label(header_labels, cd.col_idx)
                    old_val, new_val = _render_cell_pair_rich(cd.old_cell, cd.new_cell)
                    is_purple = is_subkey_row and cd.col_idx in key_cols
                    _write_row(
                        ws, row_idx,
                        [old_path, new_path, sheet_diff.name, "変更", key_val,
                         sub_key_val, item_label, old_val, new_val],
                        None,
                        subkey=is_purple,
                    )
                    row_idx += 1

    for i, width in enumerate(_COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    return wb
