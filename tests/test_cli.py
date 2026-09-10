"""
CLI引数 → DiffConfig 変換のユニットテスト。

実行:
  python tests/test_cli.py
"""
from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import argparse

from excel_diff.__main__ import (
    _build_parser, _build_config, _apply_profile, _profiles_dir,
    _save_workbook_or_raise, _write_index_xlsx,
)


# ---------------------------------------------------------------------------
# ヘルパー
# ---------------------------------------------------------------------------

PASS = []
FAIL = []


def _run_test(name: str, fn):
    try:
        fn()
        print(f"  PASS  {name}")
        PASS.append(name)
    except AssertionError as e:
        print(f"  FAIL  {name}: {e}")
        FAIL.append(name)
    except Exception as e:
        print(f"  ERROR {name}: {type(e).__name__}: {e}")
        FAIL.append(name)


def build_config(argv: list[str]):
    args = _build_parser().parse_args(argv)
    return _build_config(args)


def assert_raises(exc_type, fn):
    try:
        fn()
    except exc_type:
        return
    except Exception as e:
        raise AssertionError(
            f"{exc_type.__name__} ではなく {type(e).__name__} が発生した: {e}"
        )
    raise AssertionError(f"{exc_type.__name__} が発生しなかった")


# ---------------------------------------------------------------------------
# テストケース
# ---------------------------------------------------------------------------

def t_sub_key_cols_parsed_with_key_cols():
    """--key-cols と --sub-key-cols を併用すると、両方が DiffConfig に反映される"""
    config = build_config(["--key-cols", "A", "--sub-key-cols", "B"])
    assert config.diff_mode == "key"
    assert config.key_cols == [0]
    assert config.sub_key_cols == [1]


def t_sub_key_cols_without_key_cols_exits():
    """--key-cols なしで --sub-key-cols だけ指定するとエラー終了する"""
    assert_raises(SystemExit, lambda: build_config(["--sub-key-cols", "B"]))


def t_sub_key_cols_overlap_with_key_cols_exits():
    """--key-cols と --sub-key-cols に同じ列を指定するとエラー終了する"""
    assert_raises(
        SystemExit,
        lambda: build_config(["--key-cols", "A", "--sub-key-cols", "A"]),
    )


def t_profile_dir_diff_applies_sub_key_cols():
    """--profile（dir_diffタブ相当）の sub_key_cols が args に反映される
    （p-pipelineが --dir ... --profile <name> で使う経路）。"""
    profile_name = "__test_subkey_profile__"
    profile_path = _profiles_dir() / f"{profile_name}.toml"
    profile_path.write_text(
        '[dir_diff]\ndiff_mode = "key"\nkey_cols = "A"\nsub_key_cols = "B"\n',
        encoding="utf-8",
    )
    try:
        args = argparse.Namespace(
            profile=profile_name,
            dir=["old_dir", "new_dir"],
            split=None,
            old_file=None,
            new_file=None,
            output_dir=None, sheet_old=None, sheet_new=None,
            include_cols=None, matchers=None, strikethrough=False,
            open=True, diff_mode=None, key_cols=None, sub_key_cols=None,
        )
        _apply_profile(args)
        assert args.key_cols == "A", f"key_cols が {args.key_cols!r}"
        assert args.sub_key_cols == "B", f"sub_key_cols が {args.sub_key_cols!r}（反映されていない）"
    finally:
        profile_path.unlink(missing_ok=True)


def t_matchers_json_with_invalid_subkey_config_exits_cleanly():
    """--matchers のJSONに不正な sub_key_cols 設定があった場合、生の
    ValueError で落ちず、他のバリデーションと同様にエラー終了する。"""
    import json
    import tempfile
    data = {"diff_mode": "lcs", "sub_key_cols": "B", "matchers": []}
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".json", delete=False, encoding="utf-8"
    ) as f:
        json.dump(data, f)
        path = f.name
    try:
        assert_raises(SystemExit, lambda: build_config(["--matchers", path]))
    finally:
        os.remove(path)


def t_sub_key_cols_with_lcs_mode_exits():
    """--diff-mode lcs と --sub-key-cols を併用するとエラー終了する"""
    assert_raises(
        SystemExit,
        lambda: build_config(
            ["--key-cols", "A", "--diff-mode", "lcs", "--sub-key-cols", "B"]
        ),
    )


# ---------------------------------------------------------------------------
# --excel-summary / --header-row
# ---------------------------------------------------------------------------

def t_excel_summary_without_key_mode_exits():
    """--excel-summary は --diff-mode key（--key-cols指定）が前提。lcsモードだとエラー終了する"""
    assert_raises(SystemExit, lambda: build_config(["--excel-summary"]))


def t_excel_summary_with_key_mode_succeeds():
    """--key-cols併用（keyモード）なら --excel-summary はエラーにならない"""
    config = build_config(["--key-cols", "A", "--excel-summary"])
    assert config.diff_mode == "key"


def t_excel_summary_default_is_none():
    """--excel-summary 未指定時は args.excel_summary が None のまま"""
    args = _build_parser().parse_args(["--key-cols", "A"])
    assert args.excel_summary is None


def t_excel_summary_bare_flag_yields_empty_string():
    """--excel-summary をパス省略で指定すると const の空文字列になる"""
    args = _build_parser().parse_args(["--key-cols", "A", "--excel-summary"])
    assert args.excel_summary == ""


def t_excel_summary_with_path():
    """--excel-summary にパスを指定するとそのまま args に反映される"""
    args = _build_parser().parse_args(["--key-cols", "A", "--excel-summary", "out.xlsx"])
    assert args.excel_summary == "out.xlsx"


def t_header_row_default_is_one():
    """--header-row 未指定時のデフォルトは1"""
    args = _build_parser().parse_args([])
    assert args.header_row == 1


def t_header_row_explicit_value():
    """--header-row を明示指定すると反映される（0=ヘッダーなし扱いも許可）"""
    args = _build_parser().parse_args(["--header-row", "0"])
    assert args.header_row == 0
    args = _build_parser().parse_args(["--header-row", "5"])
    assert args.header_row == 5


def t_profile_dir_diff_applies_excel_summary_and_header_row():
    """--profile（dir_diffタブ相当）の excel_summary/header_row が args に反映される"""
    profile_name = "__test_excel_summary_profile__"
    profile_path = _profiles_dir() / f"{profile_name}.toml"
    profile_path.write_text(
        '[dir_diff]\ndiff_mode = "key"\nkey_cols = "A"\n'
        'excel_summary = true\nheader_row = 5\n',
        encoding="utf-8",
    )
    try:
        args = argparse.Namespace(
            profile=profile_name,
            dir=["old_dir", "new_dir"],
            split=None,
            old_file=None,
            new_file=None,
            output_dir=None, sheet_old=None, sheet_new=None,
            include_cols=None, matchers=None, strikethrough=False,
            open=True, diff_mode=None, key_cols=None, sub_key_cols=None,
            excel_summary=None, header_row=1,
        )
        _apply_profile(args)
        assert args.excel_summary is True, f"excel_summary が {args.excel_summary!r}"
        assert args.header_row == 5, f"header_row が {args.header_row!r}"
    finally:
        profile_path.unlink(missing_ok=True)


# ---------------------------------------------------------------------------
# Excel保存時の PermissionError（他プロセスで開いている場合）ハンドリング
# ---------------------------------------------------------------------------

def _locked_temp_xlsx_path():
    """書き込みロックした一時ファイルのパスと、ロック解除用のfileオブジェクトを返す。
    Windows専用（msvcrt.locking を使用、本ツールはWindows専用のためこれで良い）。"""
    import msvcrt
    import tempfile

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    lock_f = open(path, "wb")
    lock_f.write(b"dummy")
    lock_f.flush()
    msvcrt.locking(lock_f.fileno(), msvcrt.LK_NBLCK, 1)
    return path, lock_f


def _unlock_and_remove(path, lock_f):
    """ロック解除してファイルを削除する。

    _save_workbook_or_raise() は raise ... from first_error で元の例外を連鎖
    させるため、その中で openpyxl が生成した（クローズに失敗した）ZipFile
    オブジェクトがトレースバック経由で参照され続け、循環参照GCが走るまで
    OSファイルハンドルが解放されないことがある（本番のCLI/GUIではエラー後
    すぐプロセス終了/ログ表示のみのため実害はないが、同一プロセス内で
    すぐ削除するテストでは gc.collect() で明示的に回収してやる必要がある）。
    """
    import gc
    import msvcrt

    msvcrt.locking(lock_f.fileno(), msvcrt.LK_UNLCK, 1)
    lock_f.close()
    gc.collect()
    os.remove(path)


def t_save_workbook_or_raise_fixes_openpyxl_crlf_writing():
    """openpyxl（3.1.5で確認）はセル値中の \\n を保存時に \\r\\n として書き込んで
    しまう既知の挙動があり、生の \\r がXMLに残るとOOXMLの往復仕様上不正で、
    Excel起動時に「修復されたレコード」警告（sheet1.xml内の文字列プロパティ）の
    原因になる（実運用で発生した不具合の回帰テスト）。_save_workbook_or_raise()
    が保存後にこれを正規化し、生のXMLに \\r が残らないことを検証する。"""
    import tempfile
    import xml.etree.ElementTree as ET
    import zipfile

    import openpyxl

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        wb = openpyxl.Workbook()
        wb.active["A1"] = "line1\nline2"  # LFのみ（CRは含めない）でも発生する
        _save_workbook_or_raise(wb, path, "テスト")

        with zipfile.ZipFile(path) as z:
            for name in z.namelist():
                if not name.endswith(".xml"):
                    continue
                data = z.read(name)
                assert b"\r" not in data, f"{name} に生の\\rが残っている"
                ET.fromstring(data)  # 構文的にも正しいXMLであること

        result = openpyxl.load_workbook(path)
        assert result.active["A1"].value == "line1\nline2", (
            f"往復後の値が変化した: {result.active['A1'].value!r}"
        )
    finally:
        os.remove(path)


def t_save_workbook_or_raise_adds_xml_space_preserve():
    """`<t>` 要素の内容が空白のみ/前後が空白（例: 差分がLF1文字だけの変更run）の
    場合、openpyxl（3.1.5で確認）は xml:space="preserve" を付与しない既知の
    制限があり、Excelが「修復されたレコード」警告を出す原因になる（プレーン
    セル・リッチテキストのrun両方で発生。実運用で発生した不具合の回帰テスト）。
    _save_workbook_or_raise() が保存後にこれを補完することを検証する。"""
    import tempfile
    import xml.etree.ElementTree as ET
    import zipfile

    import openpyxl
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "\n"  # プレーンセル: 内容が空白(改行)のみ
        font = InlineFont(color="CF222E", strike=True)
        ws["A2"] = CellRichText(["prefix", TextBlock(font, "\n"), "suffix"])
        _save_workbook_or_raise(wb, path, "テスト")

        with zipfile.ZipFile(path) as z:
            data = z.read("xl/worksheets/sheet1.xml")
            ET.fromstring(data)  # 構文的に正しいXMLであること
            assert b'<t xml:space="preserve">\n</t>' in data, (
                "xml:space=preserve が付与されていない"
            )

        result = openpyxl.load_workbook(path, rich_text=True)
        rws = result.active
        assert rws["A1"].value == "\n", f"A1の往復後の値が変化した: {rws['A1'].value!r}"
        assert str(rws["A2"].value) == "prefix\nsuffix", (
            f"A2の往復後の値が変化した: {rws['A2'].value!r}"
        )
    finally:
        os.remove(path)


def t_save_workbook_or_raise_wraps_permission_error():
    """保存先が他プロセスで開かれている（PermissionError）場合、原因と対処法を
    含む分かりやすいメッセージに変換して再送出する（生のトレースバックで
    落ちてp-pipeline経由の実行がexit code 1にはなるがメッセージが分かりにくい、
    という実運用での不具合の回帰テスト）。"""
    import openpyxl

    path, lock_f = _locked_temp_xlsx_path()
    try:
        wb = openpyxl.Workbook()

        def do_save():
            _save_workbook_or_raise(wb, path, "テストExcel")

        assert_raises(PermissionError, do_save)
        try:
            do_save()
        except PermissionError as e:
            assert "テストExcel" in str(e), f"ラベルがメッセージに含まれない: {e}"
            assert "閉じてから再実行" in str(e), f"対処法がメッセージに含まれない: {e}"
    finally:
        _unlock_and_remove(path, lock_f)


def t_save_workbook_or_raise_closes_open_excel_and_retries():
    """保存先が起動中のExcelで開かれている場合、自動的に閉じて（保存確認なしで
    破棄）保存をリトライする。ユーザー要望「ファイルが開いていたら閉じてから
    再試行してほしい」に対応する回帰テスト。実際にExcelをCOM経由で起動して
    対象ファイルを開いた状態から検証する。"""
    import tempfile

    import openpyxl
    import pythoncom
    import win32com.client

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = os.path.abspath(f.name)

    wb0 = openpyxl.Workbook()
    wb0.active["A1"] = "original"
    wb0.save(path)

    pythoncom.CoInitialize()
    app = None
    try:
        app = win32com.client.Dispatch("Excel.Application")
        app.Visible = False
        app.Workbooks.Open(path)

        wb_new = openpyxl.Workbook()
        wb_new.active["A1"] = "new content"
        _save_workbook_or_raise(wb_new, path, "テストExcel")  # 例外が出ないこと

        result = openpyxl.load_workbook(path)
        assert result.active["A1"].value == "new content", (
            f"自動クローズ後の保存内容が反映されていない: {result.active['A1'].value!r}"
        )
    finally:
        try:
            if app is not None:
                app.Quit()
        except Exception:
            pass  # 最後のブックが閉じた時点でExcel自体が終了している場合がある
        pythoncom.CoUninitialize()
        os.remove(path)


def t_write_index_xlsx_permission_error_is_friendly():
    """_write_index_xlsx() 自体も、保存先が他プロセスで開かれている場合に
    分かりやすいメッセージのPermissionErrorを送出する
    （実際の障害: p-pipelineでの ★summary.xlsx 保存失敗と同根の箇所）。"""
    path, lock_f = _locked_temp_xlsx_path()
    try:
        def do_write():
            _write_index_xlsx([], [], "old_dir", "new_dir", path)

        assert_raises(PermissionError, do_write)
        try:
            do_write()
        except PermissionError as e:
            assert "インデックス" in str(e), f"ラベルがメッセージに含まれない: {e}"
            assert "閉じてから再実行" in str(e), f"対処法がメッセージに含まれない: {e}"
    finally:
        _unlock_and_remove(path, lock_f)


# ---------------------------------------------------------------------------
# メイン
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    print("=" * 50)
    print("excel-diff CLIユニットテスト")
    print("=" * 50)

    _run_test("サブキー: --key-colsと併用で反映",       t_sub_key_cols_parsed_with_key_cols)
    _run_test("サブキー: --key-cols無しでエラー終了",   t_sub_key_cols_without_key_cols_exits)
    _run_test("サブキー: --key-colsと重複でエラー終了", t_sub_key_cols_overlap_with_key_cols_exits)
    _run_test("サブキー: lcsモードでエラー終了",        t_sub_key_cols_with_lcs_mode_exits)
    _run_test("サブキー: --profile(dir_diff)で反映",    t_profile_dir_diff_applies_sub_key_cols)
    _run_test("サブキー: matchers JSON不正設定でエラー終了", t_matchers_json_with_invalid_subkey_config_exits_cleanly)

    print()
    _run_test("集約Excel: keyモード無しでエラー終了",     t_excel_summary_without_key_mode_exits)
    _run_test("集約Excel: keyモード併用で成功",           t_excel_summary_with_key_mode_succeeds)
    _run_test("集約Excel: 未指定はNone",                 t_excel_summary_default_is_none)
    _run_test("集約Excel: パス省略でconst空文字列",       t_excel_summary_bare_flag_yields_empty_string)
    _run_test("集約Excel: パス指定が反映される",           t_excel_summary_with_path)
    _run_test("ヘッダー行: デフォルトは1",                t_header_row_default_is_one)
    _run_test("ヘッダー行: 明示指定が反映される",          t_header_row_explicit_value)
    _run_test(
        "集約Excel: --profile(dir_diff)で反映",
        t_profile_dir_diff_applies_excel_summary_and_header_row,
    )

    print()
    _run_test(
        "保存エラー: openpyxlのCRLF書き込みを正規化",
        t_save_workbook_or_raise_fixes_openpyxl_crlf_writing,
    )
    _run_test(
        "保存エラー: 空白のみのrunにxml:space=preserveを付与",
        t_save_workbook_or_raise_adds_xml_space_preserve,
    )
    _run_test(
        "保存エラー: PermissionErrorを分かりやすく変換",
        t_save_workbook_or_raise_wraps_permission_error,
    )
    _run_test(
        "保存エラー: Excelで開いていれば閉じてリトライ",
        t_save_workbook_or_raise_closes_open_excel_and_retries,
    )
    _run_test(
        "保存エラー: _write_index_xlsxも分かりやすく変換",
        t_write_index_xlsx_permission_error_is_friendly,
    )

    print("=" * 50)
    print(f"結果: {len(PASS)} PASS / {len(FAIL)} FAIL")
    if FAIL:
        sys.exit(1)
